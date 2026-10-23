"""Génère les versions Word et les actifs du rapport MANUGOM Stock.

Le DOCX est la source éditable. Microsoft Word l'exporte ensuite en PDF afin
que les deux livrables conservent la même pagination et la même mise en page.
"""

from __future__ import annotations

import json
import sqlite3
import sys
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.graphics import renderPDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[2]
OUTPUT_DOCX = ROOT / "output" / "docx" / "Rapport_MANUGOM_Stock.docx"
OUTPUT_PDF = ROOT / "output" / "pdf" / "Rapport_MANUGOM_Stock.pdf"
ASSET_DIR = ROOT / "docs" / "rapport" / "assets"
SCREENSHOT_DIR = ASSET_DIR / "screenshots"
DIAGRAM_DIR = ASSET_DIR / "diagrams"
TMP_DIR = ROOT / "tmp" / "pdfs" / "report-v2"

RED = "C90000"
GREEN = "2F6B55"
TEXT = "111111"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
BORDER = "D9D9D9"

sys.path.insert(0, str(ROOT / "docs" / "uml"))
from staruml_style import (  # noqa: E402
    activity_diagram,
    architecture_diagram,
    class_diagram,
    class_supply_diagram,
    class_traceability_diagram,
    data_model_overview,
    deployment_diagram,
    sequence_confirmation_diagram,
    sequence_diagram,
    use_case_diagram,
    use_case_management_diagram,
    write_svgs,
)


FIGURE_FACTORIES = {
    "uml-cas-public-client.png": use_case_diagram,
    "uml-cas-gestion.png": use_case_management_diagram,
    "architecture-mvt.png": architecture_diagram,
    "uml-classes-catalogue-commandes.png": class_diagram,
    "uml-classes-approvisionnements.png": class_supply_diagram,
    "uml-classes-tracabilite.png": class_traceability_diagram,
    "uml-classes-comptes.png": data_model_overview,
    "uml-sequence-creation.png": sequence_diagram,
    "uml-sequence-confirmation.png": sequence_confirmation_diagram,
    "uml-activite-commande.png": activity_diagram,
    "uml-deploiement.png": deployment_diagram,
}


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=BORDER, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate])
    if placeholder:
        text_run = paragraph.add_run(placeholder)
        set_run_font(text_run, size=10, color=GRAY)
    end_run = paragraph.add_run()
    end_run._r.append(end)


def set_page_number_start(section, start=1):
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def set_table_width(cell, width_inches):
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def anchor_picture_to_page(shape):
    """Place an inline picture at the physical top-left corner of the page."""
    drawing = shape._inline
    drawing.tag = qn("wp:anchor")
    for name, value in (
        ("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"),
        ("simplePos", "0"), ("relativeHeight", "0"), ("behindDoc", "0"),
        ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1"),
    ):
        drawing.set(name, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0"); simple_pos.set("y", "0")
    position_h = OxmlElement("wp:positionH"); position_h.set("relativeFrom", "page")
    offset_h = OxmlElement("wp:posOffset"); offset_h.text = "0"; position_h.append(offset_h)
    position_v = OxmlElement("wp:positionV"); position_v.set("relativeFrom", "page")
    offset_v = OxmlElement("wp:posOffset"); offset_v.text = "0"; position_v.append(offset_v)
    wrap_none = OxmlElement("wp:wrapNone")

    drawing.insert(0, simple_pos)
    drawing.insert(1, position_h)
    drawing.insert(2, position_v)
    effect_extent = drawing.find(qn("wp:effectExtent"))
    insert_at = drawing.index(effect_extent) + 1 if effect_extent is not None else 4
    drawing.insert(insert_at, wrap_none)


def make_cover_image():
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    cover_pdf = TMP_DIR / "cover.pdf"
    cover_png = ASSET_DIR / "cover-red.png"
    font_regular = Path("C:/Windows/Fonts/times.ttf")
    font_bold = Path("C:/Windows/Fonts/timesbd.ttf")
    if font_regular.exists() and font_bold.exists():
        pdfmetrics.registerFont(TTFont("CoverTimes", str(font_regular)))
        pdfmetrics.registerFont(TTFont("CoverTimesBold", str(font_bold)))
        regular, bold = "CoverTimes", "CoverTimesBold"
    else:
        regular, bold = "Times-Roman", "Times-Bold"
    c = canvas.Canvas(str(cover_pdf), pagesize=A4)
    w, h = A4
    c.setFillColor(colors.white)
    c.rect(0, 0, w, h, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#C90000"))
    p = c.beginPath(); p.moveTo(0, h); p.lineTo(0, h-245); p.lineTo(315, h); p.close(); c.drawPath(p, fill=1, stroke=0)
    c.setStrokeColor(colors.HexColor("#C90000")); c.setLineWidth(1.2); c.line(0, h-258, 324, h)
    p = c.beginPath(); p.moveTo(w, 0); p.lineTo(w-225, 0); p.lineTo(w, 225); p.close(); c.drawPath(p, fill=1, stroke=0)
    c.setFillColor(colors.black)
    c.setFont(regular, 14); c.drawCentredString(w/2, h-175, "MANUGOM STOCK")
    c.setFont("Times-Italic", 11); c.drawCentredString(w/2, h-194, "Projet personnel de développement web")
    c.setFont(bold, 24); c.drawCentredString(w/2, h-350, "Application Web de Gestion de Stock")
    c.drawCentredString(w/2, h-380, "et de Commandes de Pneus")
    c.setFont("Times-Italic", 10); c.setFillColor(colors.HexColor("#6A6A6A"))
    c.drawCentredString(w/2, h-420, "Rapport de conception, de réalisation et de validation")
    fields = [("Réalisé par :", "Youssef Bouikourar")]
    y = 285
    for label, value in fields:
        c.setFillColor(colors.black); c.setFont(bold, 10); c.drawString(104, y, label)
        c.setFont(regular, 10); c.drawString(255, y, value); y -= 31
    c.setFillColor(colors.HexColor("#666666")); c.setFont(regular, 9)
    c.showPage(); c.save()
    doc = pdfium.PdfDocument(str(cover_pdf))
    doc[0].render(scale=3.0).to_pil().save(cover_png)
    return cover_png


def render_diagram_assets():
    DIAGRAM_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    write_svgs(ROOT / "docs" / "uml")
    result = {}
    for filename, factory in FIGURE_FACTORIES.items():
        drawing = factory()
        pdf_path = TMP_DIR / filename.replace(".png", ".pdf")
        png_path = DIAGRAM_DIR / filename
        renderPDF.drawToFile(drawing, str(pdf_path))
        doc = pdfium.PdfDocument(str(pdf_path))
        doc[0].render(scale=2.6).to_pil().save(png_path)
        result[filename] = png_path
    return result


class ReportBuilder:
    def __init__(self, document, diagrams):
        self.doc = document
        self.diagrams = diagrams
        self.figure_number = 0
        self.table_number = 0

    def p(self, text="", *, style=None, align=None, bold_lead=None, keep_with_next=False):
        p = self.doc.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        p.paragraph_format.keep_with_next = keep_with_next
        if bold_lead and text.startswith(bold_lead):
            r1 = p.add_run(bold_lead); set_run_font(r1, bold=True)
            r2 = p.add_run(text[len(bold_lead):]); set_run_font(r2)
        else:
            r = p.add_run(text); set_run_font(r)
        return p

    def h1(self, title):
        p = self.doc.add_paragraph(title, style="Heading 1")
        return p

    def h2(self, title):
        p = self.doc.add_paragraph(title, style="Heading 2")
        p.paragraph_format.keep_with_next = True
        return p

    def h3(self, title):
        return self.doc.add_paragraph(title, style="Heading 3")

    def bullets(self, items):
        for item in items:
            p = self.doc.add_paragraph(style="List Bullet")
            r = p.add_run(item); set_run_font(r)

    def code(self, lines):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.right_indent = Cm(0.5)
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(8)
        p._p.get_or_add_pPr().append(parse_xml(r'<w:shd {} w:fill="F3F3F3"/>'.format(nsdecls("w"))))
        for index, line in enumerate(lines):
            r = p.add_run(line); set_run_font(r, "Courier New", 9.5)
            if index < len(lines)-1:
                r.add_break()

    def table(self, title, headers, rows, widths):
        self.table_number += 1
        caption = self.doc.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = caption.add_run("Tableau "); set_run_font(r, size=10, italic=True)
        add_field(caption, " SEQ Tableau \\* ARABIC ", str(self.table_number))
        r = caption.add_run(" : " + title); set_run_font(r, size=10, italic=True)
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_repeat_table_header(table.rows[0])
        for col, (cell, value) in enumerate(zip(table.rows[0].cells, headers)):
            set_table_width(cell, widths[col]); set_cell_shading(cell, LIGHT_GRAY); set_cell_borders(cell); set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value)); set_run_font(r, "Arial", 9.5, bold=True)
        for row_index, values in enumerate(rows):
            row = table.add_row(); set_no_split(row)
            for col, (cell, value) in enumerate(zip(row.cells, values)):
                set_table_width(cell, widths[col]); set_cell_borders(cell); set_cell_margins(cell)
                if row_index % 2 == 1:
                    set_cell_shading(cell, "FAFAFA")
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col == 0 and len(widths) > 2 else WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(str(value)); set_run_font(r, size=9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(1)
        return table

    def figure(self, image_path, title, explanation, *, width=6.25, page_before=False):
        if page_before:
            previous = self.doc.paragraphs[-1] if self.doc.paragraphs else None
            if previous is not None and previous.style.name.startswith("Heading"):
                previous.paragraph_format.page_break_before = True
                previous.paragraph_format.keep_with_next = True
            else:
                self.doc.add_page_break()
        picture = self.doc.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.keep_with_next = True
        shape = picture.add_run().add_picture(str(image_path), width=Inches(width))
        shape._inline.docPr.set("title", title)
        shape._inline.docPr.set("descr", explanation)
        self.figure_number += 1
        caption = self.doc.add_paragraph(style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.keep_with_next = True
        r = caption.add_run("Figure "); set_run_font(r, size=10, italic=True)
        add_field(caption, " SEQ Figure \\* ARABIC ", str(self.figure_number))
        r = caption.add_run(" : " + title); set_run_font(r, size=10, italic=True)
        self.p("Interprétation de la figure. " + explanation, bold_lead="Interprétation de la figure. ")


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(12); normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    pf.widow_control = True
    for name, size, before, after in (("Heading 1", 14, 10, 12), ("Heading 2", 12, 10, 5), ("Heading 3", 12, 7, 4)):
        style = doc.styles[name]
        style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = True; style.font.color.rgb = RGBColor(0,0,0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.page_break_before = name == "Heading 1"
        p_bdr = style._element.pPr.find(qn("w:pBdr")) if style._element.pPr is not None else None
        if p_bdr is not None:
            style._element.pPr.remove(p_bdr)
    caption = doc.styles["Caption"]
    caption.font.name = "Times New Roman"; caption.font.size = Pt(10); caption.font.italic = True; caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman"); caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.paragraph_format.space_before = Pt(4); caption.paragraph_format.space_after = Pt(6)
    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"; style.font.size = Pt(12)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.space_after = Pt(2)


def configure_section(section, *, margins=True, header_text=None, page_number=False, restart=None):
    section.page_width = Cm(21); section.page_height = Cm(29.7); section.orientation = WD_ORIENT.PORTRAIT
    margin = Cm(2.5) if margins else Cm(0)
    section.top_margin = margin; section.bottom_margin = margin; section.left_margin = margin; section.right_margin = margin
    section.header_distance = Cm(1); section.footer_distance = Cm(1)
    section.header.is_linked_to_previous = False; section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]; hp.clear(); hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if header_text:
        r = hp.add_run(header_text); set_run_font(r, "Arial", 8, color=GREEN)
    fp = section.footer.paragraphs[0]; fp.clear(); fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if page_number:
        add_field(fp, " PAGE ", "1")
    if restart is not None:
        set_page_number_start(section, restart)
    else:
        # A new python-docx section inherits the previous section properties.
        # Drop an inherited explicit start so numbering continues between
        # chapters instead of restarting at page 1.
        pg_num = section._sectPr.find(qn("w:pgNumType"))
        if pg_num is not None and qn("w:start") in pg_num.attrib:
            del pg_num.attrib[qn("w:start")]


def new_body_section(doc, header, restart=None):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, margins=True, header_text=header, page_number=True, restart=restart)
    return section


def add_front_heading(builder, title):
    p = builder.doc.add_paragraph(title, style="Heading 1")
    p.paragraph_format.page_break_before = False


def read_stats():
    db = ROOT / "db.sqlite3"
    tables = {
        "Utilisateurs": "accounts_user", "Clients": "customers_client", "Marques": "inventory_marque",
        "Pneus actifs": "inventory_pneu", "Commandes": "orders_commande",
        "Approvisionnements": "suppliers_approvisionnement", "Mouvements": "inventory_mouvementstock",
    }
    connection = sqlite3.connect(f"file:{db.as_posix()}?mode=ro", uri=True)
    result = {name: connection.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0] for name, table in tables.items()}
    connection.close()
    return result


def build_docx():
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    diagrams = render_diagram_assets()
    cover = make_cover_image()
    stats = read_stats()
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], margins=False)
    settings = doc.settings._element
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true"); settings.append(update)
    doc.core_properties.title = "Application Web de Gestion de Stock et de Commandes de Pneus"
    doc.core_properties.subject = "Rapport du projet personnel MANUGOM Stock"
    doc.core_properties.author = "Youssef Bouikourar"
    builder = ReportBuilder(doc, diagrams)

    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = Pt(1)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_shape = p.add_run().add_picture(str(cover), width=Cm(21), height=Cm(29.7))
    anchor_picture_to_page(cover_shape)
    cover_shape._inline.docPr.set("title", "Couverture du rapport MANUGOM Stock")
    cover_shape._inline.docPr.set("descr", "Couverture blanche et rouge du projet MANUGOM Stock, réalisée par Youssef Bouikourar.")

    front = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(front, margins=True)
    add_front_heading(builder, "Remerciements")
    builder.p("Je remercie les personnes qui ont échangé avec moi sur ce projet et partagé leurs remarques. Ces échanges m’ont aidé à préciser les besoins, à vérifier les parcours et à améliorer la présentation de l’application.")
    builder.p("Je remercie également les auteurs des documentations et les communautés des outils utilisés. Leurs ressources m’ont permis d’approfondir mes connaissances en développement web, en bases de données et en modélisation, puis de les appliquer à MANUGOM Stock.")
    builder.p("Je remercie enfin toutes les personnes qui ont contribué, directement ou indirectement, à l'aboutissement de ce travail. Les phases de vérification, les échanges autour des besoins et les essais de l'application ont permis de corriger plusieurs incohérences et de produire un prototype plus clair, documenté et démontrable.")

    doc.add_page_break(); add_front_heading(builder, "Résumé")
    builder.p("Le présent rapport décrit la conception et la réalisation de MANUGOM Stock, une application web destinée à centraliser un catalogue de pneumatiques, les commandes clients, les approvisionnements et les mouvements de stock. Le projet répond à un besoin de cohérence entre les références commerciales, les quantités disponibles et l'historique des opérations. Je l'ai développé comme projet personnel pour mettre en pratique la conception et la réalisation d'une application web complète.")
    builder.p("La solution repose sur Django 6.1.1, une base SQLite unique, des templates HTML, Bootstrap et JavaScript. Elle distingue le visiteur, le client, l'employé et l'administrateur. Le visiteur consulte le catalogue et transmet une demande de devis. Le client gère son profil, son panier et ses commandes. L'employé traite les commandes et les stocks. L'administrateur dispose en plus de la gestion des utilisateurs et de l'administration Django.")
    builder.p("Les opérations sensibles sont regroupées dans des services métier. Une commande reste en attente jusqu'à sa confirmation ; la disponibilité est alors vérifiée une seconde fois avant la création des sorties de stock. La réception d'un approvisionnement commandé produit les entrées correspondantes. Chaque mouvement conserve le stock avant et après l'opération. La version documentée a passé 34 tests automatisés et un contrôle d'intégrité SQLite. Ces résultats valident les scénarios couverts, sans constituer un essai de charge ni une preuve d'absence totale d'anomalies.")
    builder.p("Mots-clés : Django, SQLite, gestion de stock, pneumatiques, commandes, approvisionnements, UML, contrôle d'accès, tests logiciels.")

    doc.add_page_break(); add_front_heading(builder, "Abstract")
    builder.p("This report presents the design and implementation of MANUGOM Stock, a web application that centralises a tyre catalogue, customer orders, procurement records and stock movements. The project addresses the need to keep commercial references, available quantities and operation history consistent. I developed it as a personal project to practise designing, implementing and testing a complete web application.")
    builder.p("The solution uses Django 6.1.1, a single SQLite database, server-rendered HTML templates, Bootstrap and JavaScript. It distinguishes public visitors, customers, employees and administrators. Visitors browse the catalogue and submit quotation requests. Customers maintain a profile, cart and order history. Employees process orders and stock operations, while administrators also manage users and the technical Django administration area.")
    builder.p("Sensitive operations are implemented in business services. An order remains pending until an employee confirms it; stock is checked again before quantities are removed. Receiving an ordered procurement record creates the corresponding stock entries. Each movement records quantities before and after the operation. The documented version passed 34 automated tests and an SQLite integrity check. These results validate the executed scenarios but do not constitute a load test or a proof that the system is defect-free.")
    builder.p("Keywords: Django, SQLite, inventory management, tyres, orders, procurement, UML, access control, software testing.")

    doc.add_page_break(); add_front_heading(builder, "Table des matières")
    toc = doc.add_paragraph(); add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ', "Mettre à jour le champ dans Word")
    doc.add_page_break(); add_front_heading(builder, "Liste des figures")
    lof = doc.add_paragraph(); add_field(lof, ' TOC \\h \\z \\c "Figure" ', "Mettre à jour le champ dans Word")
    doc.add_page_break(); add_front_heading(builder, "Liste des tableaux")
    lot = doc.add_paragraph(); add_field(lot, ' TOC \\h \\z \\c "Tableau" ', "Mettre à jour le champ dans Word")
    doc.add_page_break(); add_front_heading(builder, "Liste des abréviations")
    builder.table("Abréviations utilisées", ["Abréviation", "Signification"], [
        ["ACID", "Atomicité, cohérence, isolation et durabilité"], ["CRUD", "Créer, consulter, modifier et supprimer"],
        ["CSRF", "Cross-Site Request Forgery"],
        ["FK", "Clé étrangère"], ["HTTP", "Hypertext Transfer Protocol"], ["IIR", "Ingénierie Informatique et Réseaux"],
        ["MCD", "Modèle conceptuel de données"], ["MLD", "Modèle logique de données"], ["MVT", "Model View Template"],
        ["ORM", "Object Relational Mapping"], ["RBAC", "Role-Based Access Control"],
        ["UML", "Unified Modeling Language"],
    ], [1.35, 4.9])

    new_body_section(doc, "Introduction générale", restart=1)
    builder.h1("Introduction générale")
    builder.p("La gestion d'un stock de pneumatiques ne se limite pas à afficher une quantité. Une référence associe une marque, un modèle, une dimension, des indices techniques, une saison, un prix et un état commercial. Une commande ajoute une autre contrainte : le système doit conserver ce que le client a commandé tout en tenant compte du stock qui peut évoluer entre la constitution du panier et le traitement par le magasin.")
    builder.p("MANUGOM Stock a été conçu pour réunir ces informations dans une application web unique. Le catalogue est accessible au public, tandis que les opérations de commande et de gestion sont protégées par l'authentification et les rôles. Le projet cherche surtout à rendre les opérations explicables. Chaque entrée, sortie, retour ou ajustement produit une trace associée à un utilisateur et, lorsque c'est pertinent, à une commande ou à un approvisionnement.")
    builder.p("La démarche a commencé par l'analyse des acteurs et des règles de gestion. Les modèles de données et les diagrammes UML ont ensuite servi à préciser les relations et les principaux scénarios. Le développement Django a été organisé par domaines fonctionnels, avec des services dédiés aux transitions sensibles. Enfin, les tests de modèles, de services, de vues et de permissions ont été exécutés sur une base temporaire.")
    builder.p("Le rapport suit cette progression. Le premier chapitre présente le contexte et la démarche du projet. Le deuxième définit le projet, ses objectifs et son cahier des charges. Le troisième expose l'analyse, l'architecture et la conception UML. Le quatrième décrit la réalisation à l'aide de captures réelles de l'application. Le cinquième présente les tests, les résultats et les limites. La bibliographie réunit des articles scientifiques et les documentations techniques utilisées.")

    new_body_section(doc, "Chapitre 1 - Contexte et démarche du projet")
    builder.h1("Chapitre 1 - Contexte et démarche du projet")
    builder.h2("1.1 Présentation du projet personnel")
    builder.p("J'ai réalisé MANUGOM Stock pour mettre en pratique le développement d'une application web de gestion. Le projet part d'un besoin concret : réunir le catalogue de pneus, les commandes et les mouvements de stock dans un même outil. Ce rapport présente les choix de conception, les fonctionnalités réalisées et les résultats des vérifications.")
    builder.p("Le travail relie l'analyse du besoin, la modélisation des données, le développement des interfaces et la validation des règles métier. La documentation permet de comprendre le fonctionnement de l'application et de la lancer dans un environnement local.")
    builder.h2("1.2 Démarche technique")
    builder.p("La réalisation associe la conception logicielle à la configuration de l'environnement d'exécution. Django structure le traitement applicatif, SQLite conserve les données et HTTP relie le navigateur au serveur. Les mécanismes de session et de protection CSRF contrôlent les actions des utilisateurs.")
    builder.p("La modélisation UML a servi à clarifier les responsabilités avant de décrire le code. Les cas d'utilisation montrent ce que chaque acteur peut demander au système. Les diagrammes de classes présentent les entités persistées et leurs relations. Les séquences et l'activité détaillent le moment où une commande devient confirmée et produit des mouvements de stock.")
    builder.h2("1.3 Présentation du contexte MANUGOM")
    builder.p("MANUGOM SARL est présentée dans le projet comme une entreprise d'importation, de distribution et de vente de pneus et de jantes à Casablanca. Le catalogue de démonstration contient plusieurs marques, catégories et dimensions afin de représenter la diversité des références que doit traiter un magasin spécialisé. Les prix et stocks chargés par le script seed_demo sont des données de démonstration, clairement signalées dans l'interface.")
    builder.p("Le point de départ est un besoin de centralisation. Sans un outil commun, le catalogue, la disponibilité, les commandes et les réceptions fournisseurs peuvent être suivis dans des supports différents. Cette dispersion complique la recherche d'une référence, le contrôle d'une quantité et l'explication d'un écart de stock. L'application rassemble ces opérations tout en conservant un historique exploitable.")
    builder.h2("1.4 Gestion de stock et apport de la recherche")
    builder.p("La gestion des stocks est un domaine ancien de la recherche opérationnelle. Harris [6] a formalisé le compromis entre la taille des lots et les coûts associés au stockage et au lancement d'une opération. MANUGOM Stock n'implémente pas de calcul de quantité économique ni de prévision de demande. Il fournit les données élémentaires nécessaires à une future étude : quantités, seuils, entrées, sorties, dates et références concernées.")
    builder.p("Dans la version actuelle, l'aide au réapprovisionnement repose sur un seuil minimum. Le tableau de bord signale les stocks faibles et les ruptures, mais la décision d'achat appartient encore à l'employé ou à l'administrateur. Cette limite est cohérente avec le périmètre de ce projet et évite de présenter une règle simple comme un modèle d'optimisation scientifique.")

    new_body_section(doc, "Chapitre 2 - Présentation du projet")
    builder.h1("Chapitre 2 - Présentation du projet")
    builder.h2("2.1 Problématique")
    builder.p("Comment construire une application qui permette de consulter précisément les pneumatiques disponibles, de transformer un panier en commande et de mettre à jour le stock sans perdre l'historique des opérations ? Cette question implique trois exigences liées. Les références doivent être assez structurées pour être recherchées. Les rôles doivent limiter les actions. Les transitions de commande et d'approvisionnement doivent rester cohérentes avec les mouvements de stock.")
    builder.h2("2.2 Objectifs")
    builder.p("L'objectif général est de proposer un prototype web démontrable qui centralise le catalogue, les clients, les commandes, les fournisseurs, les approvisionnements et les mouvements. Le projet vise également à rendre la logique vérifiable par des tests et à fournir une documentation de lancement utilisable sur un poste Windows.")
    builder.bullets(["Permettre une recherche par référence, marque, catégorie, dimension, saison, prix et disponibilité.", "Donner au client un panier persistant et un historique limité à ses propres commandes.", "Fournir aux gestionnaires les écrans nécessaires au traitement des commandes et des stocks.", "Tracer chaque variation avec le stock avant et après l'opération.", "Conserver une base SQLite unique et des commandes d'installation simples."])
    builder.h2("2.3 Acteurs et responsabilités")
    builder.table("Acteurs du système", ["Acteur", "Responsabilités principales", "Restrictions"], [
        ["Visiteur", "Consulter le site, rechercher un pneu, demander un devis et envoyer un message.", "Aucun panier ni commande."],
        ["Client", "Gérer son profil, son panier et ses commandes.", "Accès à ses propres commandes seulement."],
        ["Employé", "Gérer les références, commandes, approvisionnements, mouvements et alertes.", "Aucune gestion des comptes administrateurs."],
        ["Administrateur", "Superviser l'activité, gérer les utilisateurs et accéder à l'administration Django.", "Actions critiques toujours contrôlées côté serveur."],
    ], [1.05, 3.15, 2.1])
    builder.h2("2.4 Cahier des charges fonctionnel")
    builder.table("Fonctionnalités attendues", ["Domaine", "Fonctions"], [
        ["Catalogue", "Recherche, filtres, tri, pagination, détail et disponibilité."],
        ["Compte client", "Inscription, connexion, profil et isolation des données."],
        ["Panier", "Ajout, modification, retrait, contrôle des quantités et calcul du total."],
        ["Commandes", "Création en attente, confirmation, préparation, expédition, livraison et annulation autorisée."],
        ["Stock", "Entrées, sorties, retours, ajustements, seuils et historique."],
        ["Approvisionnements", "Fournisseur, lignes, statut commandé et réception unique."],
        ["Pilotage", "Indicateurs, graphiques, devis, contacts et gestion des utilisateurs."],
    ], [1.35, 4.95])
    builder.h2("2.5 Exigences non fonctionnelles")
    builder.table("Exigences de qualité", ["Exigence", "Réponse du projet", "Limite"], [
        ["Sécurité", "Sessions Django, CSRF, validation serveur, rôles et filtrage par propriétaire.", "Configuration locale de démonstration."],
        ["Intégrité", "Transactions courtes, contraintes d'unicité, stock non négatif et traces métier.", "Pas d'essai multiconnexion."],
        ["Utilisabilité", "Interface responsive, navigation séparée par rôle et messages de résultat.", "Pas d'étude utilisateur formelle."],
        ["Maintenabilité", "Découpage en applications et services métier testés.", "Prototype local sans CI distante."],
        ["Portabilité", "SQLite, requirements.txt, migrations et seed_demo.", "Serveur de développement local."],
    ], [1.05, 3.35, 1.95])
    builder.h2("2.6 Règles de gestion")
    builder.bullets(["Le panier ne réserve pas la quantité disponible.", "La création produit une commande EN_ATTENTE et désactive le panier utilisé.", "La confirmation revérifie le stock avant toute sortie.", "Une annulation après confirmation restaure les quantités une seule fois si la préparation n'a pas commencé.", "Seul un approvisionnement au statut COMMANDE et non encore reçu peut être réceptionné.", "Une ligne de commande conserve la référence, la dimension et le prix observés lors de la vente."])
    builder.h2("2.7 Organisation du travail")
    builder.p("Le développement a suivi des itérations courtes : initialisation de la base et des rôles, création du catalogue, mise en place du panier et des commandes, ajout des approvisionnements, construction du tableau de bord, puis tests et documentation. Les corrections ont été intégrées au même projet afin que le code, les captures et le rapport décrivent une seule version.")

    new_body_section(doc, "Chapitre 3 - Analyse et conception")
    builder.h1("Chapitre 3 - Analyse et conception")
    builder.h2("3.1 Méthode de conception")
    builder.p("L'analyse part des opérations observables puis descend vers les données et les services. Le découpage évite un diagramme global illisible. Deux vues de cas d'utilisation séparent le parcours public et client de l'espace de gestion. Les classes sont réparties par domaine. Les séquences distinguent la création d'une commande de sa confirmation, car le stock n'est retiré qu'au second moment.")
    builder.p("Parnas [2] montre que la valeur d'une décomposition dépend des responsabilités confiées à chaque module. Ce principe éclaire ici la séparation entre comptes, inventaire, clients, commandes, fournisseurs et tableau de bord. Les services ServiceCommande, ServiceStock et ServiceApprovisionnement regroupent les changements d'état sensibles, tandis que les vues orchestrent les requêtes HTTP et rendent les templates.")
    builder.h2("3.2 Cas d'utilisation du visiteur et du client")
    builder.figure(diagrams["uml-cas-public-client.png"], "Cas d'utilisation du visiteur et du client", "Utilisateur du site regroupe les fonctions publiques : rechercher dans le catalogue et contacter le magasin ou demander un devis. Visiteur non connecté et Client connecté spécialisent cet acteur et héritent de ces fonctions. Les triangles blancs pointent vers l'acteur général. La création de compte et la connexion appartiennent au visiteur ; le client dispose du profil, du panier et de ses commandes. La relation include impose la vérification du panier et de la disponibilité lors du passage d'une commande. L'annulation reste soumise au statut de la commande.", page_before=True)
    builder.h2("3.3 Cas d'utilisation de la gestion")
    builder.figure(diagrams["uml-cas-gestion.png"], "Cas d'utilisation de l'employé et de l'administrateur", "L'employé couvre le fonctionnement quotidien du magasin. L'administrateur hérite de ces possibilités et ajoute la gestion des utilisateurs et l'accès technique à Django Admin. Les associations sont individuelles et ne traversent pas les cas d'utilisation, ce qui permet de lire la figure sans ambiguïté.", page_before=True)
    builder.h2("3.4 Architecture applicative")
    builder.figure(diagrams["architecture-mvt.png"], "Architecture Django MVT de MANUGOM Stock", "La requête HTTP est dirigée par le routage vers une vue. La vue contrôle la session, les permissions et les formulaires. Elle peut appeler un service pour une mutation sensible ou interroger l'ORM pour une lecture. Le template reçoit un contexte et produit la réponse HTML. Les templates et le navigateur n'accèdent jamais directement à SQLite.", page_before=True)
    builder.p("Le diagramme corrige une simplification fréquente : toutes les lectures ne passent pas par un service. Les pages de catalogue et le tableau de bord utilisent aussi l'ORM depuis les vues. Les services forment une couche volontaire pour les règles critiques, mais Django demeure une application MVT rendue côté serveur.")
    builder.h2("3.5 Classes du catalogue et des commandes")
    builder.figure(diagrams["uml-classes-catalogue-commandes.png"], "Classes du catalogue, du panier et des commandes", "Client possède ses paniers et commandes. Les compositions indiquent que les lignes dépendent de leur panier ou commande. Pneu est relié à une marque et une catégorie. LigneCommande contient des instantanés de libellé et de prix afin qu'une modification future du catalogue ne réécrive pas l'historique.", page_before=True)
    builder.h2("3.6 Classes des approvisionnements")
    builder.figure(diagrams["uml-classes-approvisionnements.png"], "Classes des fournisseurs et approvisionnements", "Un fournisseur peut être lié à plusieurs marques et approvisionnements. Chaque approvisionnement contient ses lignes ; chaque ligne désigne un pneu et un prix d'achat. Le responsable est un utilisateur. La réception est réservée au statut COMMANDE et ne peut être répétée.", page_before=True)
    builder.h2("3.7 Classes de traçabilité")
    builder.figure(diagrams["uml-classes-tracabilite.png"], "Classes de traçabilité du stock", "Tout mouvement concerne un pneu et un auteur. La commande et l'approvisionnement sont facultatifs afin d'autoriser les ajustements manuels. Les champs stock_avant et stock_apres fournissent une trace compréhensible sans recalculer tout l'historique.", page_before=True)
    builder.h2("3.8 Comptes et profil commercial")
    builder.figure(diagrams["uml-classes-comptes.png"], "Classes des comptes, profils et clients", "User porte l'identité de connexion et le rôle. ProfilUtilisateur conserve les coordonnées associées au compte. Client représente la fiche commerciale ; elle peut exister sans compte lorsqu'elle est créée par un gestionnaire. La relation facultative évite de forcer une authentification pour toute fiche client.", page_before=True)
    builder.h2("3.9 Modèle relationnel")
    builder.p("Le modèle relationnel de Codd [1] fournit le cadre conceptuel utilisé pour séparer les entités et exprimer leurs dépendances au moyen de clés. Le schéma de MANUGOM Stock conserve toutefois une redondance intentionnelle dans LigneCommande : la référence, le modèle, la dimension et le prix sont copiés au moment de la commande pour préserver l'historique.")
    builder.table("Principales tables et clés", ["Table", "Clé primaire", "Clés étrangères principales", "Contraintes"], [
        ["inventory_pneu", "id", "marque_id, categorie_id", "reference unique, stock >= 0"],
        ["orders_panier", "id", "client_id", "un panier actif par client"],
        ["orders_lignepanier", "id", "panier_id, pneu_id", "paire panier et pneu unique"],
        ["orders_commande", "id", "client_id", "numero unique"],
        ["orders_lignecommande", "id", "commande_id, pneu_id", "instantanés de vente"],
        ["inventory_mouvementstock", "id", "pneu, user, commande ou approvisionnement", "quantité contrôlée"],
        ["suppliers_approvisionnement", "id", "fournisseur, responsable", "numero unique"],
    ], [1.45, 0.75, 2.55, 1.55])
    builder.h2("3.10 Séquence de création d'une commande")
    builder.figure(diagrams["uml-sequence-creation.png"], "Séquence de création depuis le panier", "Le client envoie un formulaire POST. La vue délègue au service, qui relit les lignes et crée la commande et ses lignes dans une transaction. Le panier devient inactif. Aucun mouvement de sortie n'est encore produit ; l'échec d'une condition annule la transaction et renvoie un message.", page_before=True)
    builder.h2("3.11 Séquence de confirmation")
    builder.figure(diagrams["uml-sequence-confirmation.png"], "Séquence de confirmation par un gestionnaire", "L'employé ou l'administrateur demande la confirmation. Le service vérifie l'état et le stock, applique une sortie pour chaque ligne et marque la commande confirmée. Une exception provoque le rollback du groupe. Sous SQLite, select_for_update n'ajoute pas de verrou de ligne ; les tests ne simulent donc pas une forte concurrence.", page_before=True)
    builder.p("Härder et Reuter [3] décrivent les propriétés ACID des transactions. Dans ce projet, transaction.atomic regroupe les écritures qui doivent réussir ensemble. Cette utilisation protège contre une commande partiellement confirmée à la suite d'une exception. Elle ne constitue pas, à elle seule, une validation du comportement sous charge concurrente.")
    builder.h2("3.12 Activité de traitement")
    builder.figure(diagrams["uml-activite-commande.png"], "Activité du panier à la confirmation", "Le premier contrôle appartient au parcours client et crée une commande en attente. Le second appartient au gestionnaire et précède le retrait du stock. Les branches d'erreur se terminent sans modification partielle. Cette séparation explique pourquoi l'affichage du panier ne réserve pas une quantité.", page_before=True)
    builder.h2("3.13 Déploiement local")
    builder.figure(diagrams["uml-deploiement.png"], "Déploiement de démonstration", "Le navigateur communique en HTTP avec le serveur Django local. Le code, SQLite et le dossier media résident sur la machine de développement. Cette topologie correspond à la démonstration ; une production demanderait un serveur applicatif, HTTPS, des secrets adaptés, un stockage sauvegardé et une base choisie selon la concurrence attendue.", page_before=True)

    new_body_section(doc, "Chapitre 4 - Réalisation et développement")
    builder.h1("Chapitre 4 - Réalisation et développement")
    builder.h2("4.1 Environnement de développement")
    builder.table("Environnement technique", ["Élément", "Version ou rôle"], [["Python", "3.14.6 dans l'environnement local contrôlé"], ["Django", "6.1.1"], ["Pillow", "12.3.0"], ["SQLite", "3.50.4 observé avec Python"], ["Bootstrap", "Mise en page responsive"], ["Chart.js", "Graphiques du tableau de bord"], ["ReportLab et python-docx", "Documentation et diagrammes"]], [1.65, 4.65])
    builder.p("Les commandes principales restent volontairement simples. Elles créent l'environnement, installent les dépendances, préparent la base, chargent le jeu de démonstration et démarrent le serveur local.")
    builder.code(['python -m venv .venv', r'.\.venv\Scripts\Activate.ps1', 'python -m pip install -r requirements.txt', 'python manage.py migrate', 'python manage.py seed_demo', 'python manage.py runserver'])
    builder.h2("4.2 Organisation du code")
    builder.p("Le projet contient huit applications Django. accounts gère le compte personnalisé et les rôles. inventory contient les pneus, marques, catégories, mouvements et ServiceStock. customers représente les fiches clients. orders regroupe le panier, la commande et ServiceCommande. suppliers couvre les fournisseurs, les lignes d'approvisionnement et ServiceApprovisionnement. dashboard fournit l'espace de gestion. quotes enregistre les devis et contacts. core réunit les pages institutionnelles et la commande seed_demo.")
    builder.p("Les URL sont regroupées par espace de noms. Les templates réutilisent une base commune et des composants de présentation. Les fichiers statiques sont centralisés dans static, tandis que les images téléversées se trouvent dans media. Ce découpage limite le couplage entre domaines et rend les tests plus ciblés.")
    builder.h2("4.3 Authentification et contrôle des rôles")
    builder.p("Le modèle User étend AbstractUser avec trois valeurs : ADMIN, EMPLOYEE et CLIENT. Les décorateurs manager_required et admin_required protègent les vues de gestion. Les détails de commande ajoutent un contrôle de propriétaire pour qu'un client ne puisse consulter qu'une commande reliée à sa propre fiche.")
    builder.p("Sandhu, Ferraiolo et Kuhn [4] associent les autorisations aux rôles organisationnels dans le modèle RBAC. MANUGOM Stock applique une version simplifiée de cette idée. Il ne fournit pas un moteur RBAC configurable, une hiérarchie dynamique ni des contraintes formelles de séparation des tâches. Le visiteur est simplement un utilisateur anonyme.")
    builder.h2("4.4 Interfaces publiques")
    builder.figure(SCREENSHOT_DIR / "01-visiteur-accueil.png", "Page d'accueil publique", "La page présente l'activité, les coordonnées et quatre actions directes : catalogue, devis, appel et WhatsApp. Les liens téléphoniques utilisent tel: et wa.me. Ils ouvrent l'application correspondante sans appeler une API depuis Django.", page_before=True)
    builder.figure(SCREENSHOT_DIR / "02-visiteur-catalogue.png", "Catalogue et filtres de recherche", "Le visiteur recherche une référence sans connexion. Les filtres portent sur les dimensions, la marque, la catégorie, la saison, le prix et la disponibilité. Les quantités proviennent de SQLite et les cartes affichent explicitement des données de démonstration.", page_before=True)
    builder.figure(SCREENSHOT_DIR / "03-visiteur-fiche-pneu.png", "Fiche détaillée d'un pneu", "La fiche réunit les caractéristiques utiles avant une commande. Lorsqu'un visiteur souhaite ajouter l'article, le parcours lui propose de se connecter puis le ramène vers une page GET valide. Cette correction évite la réponse HTTP 405 observée lors d'une redirection vers une URL POST.", page_before=True)
    builder.h2("4.5 Espace client")
    builder.figure(SCREENSHOT_DIR / "04-client-commandes.png", "Historique des commandes du client", "Le compte client1 ne voit que ses propres commandes. Le numéro, la date, le mode de réception, le statut et le total permettent de suivre chaque demande sans exposer celles des autres clients.", page_before=True)
    builder.figure(SCREENSHOT_DIR / "05-client-detail-commande.png", "Détail d'une commande client", "Les lignes affichées reprennent les instantanés enregistrés lors de la création. Cette présentation reste stable même si un pneu change ensuite de prix ou de libellé dans le catalogue.", page_before=True)
    builder.h2("4.6 Espace employé")
    builder.figure(SCREENSHOT_DIR / "06-employe-commandes.png", "Suivi des commandes par l'employé", "La liste de gestion regroupe les commandes de tous les clients avec leur état et une action de traitement. Les transitions proposées par la fiche dépendent du statut courant et sont revérifiées côté serveur.", page_before=True)
    builder.figure(SCREENSHOT_DIR / "07-employe-stock.png", "Gestion du référentiel des pneus", "L'employé consulte les stocks, les prix et les références depuis une vue compacte. Les formulaires contrôlent la dimension et les montants. L'archivage est privilégié lorsqu'une référence possède déjà un historique.", page_before=True)
    builder.h2("4.7 Tableau de bord")
    builder.figure(SCREENSHOT_DIR / "08-dashboard-synthese.png", "Indicateurs et ventes du tableau de bord", "Les cartes sont calculées à partir de SQLite : références, quantité, valeur estimée, clients, commandes, chiffre d'affaires livré, stocks faibles et ruptures. Les graphiques présentent les ventes mensuelles et la répartition par statut.", page_before=True)
    builder.figure(SCREENSHOT_DIR / "09-dashboard-graphiques.png", "Analyses complémentaires du stock", "La suite du tableau de bord met en relation les marques, les produits vendus et les mouvements. Il s'agit d'une lecture descriptive du jeu de démonstration. Aucun modèle prédictif ou calcul de réapprovisionnement optimal n'est appliqué.", page_before=True)
    builder.h2("4.8 Administration")
    builder.figure(SCREENSHOT_DIR / "10-admin-utilisateurs.png", "Gestion des utilisateurs et des rôles", "L'administrateur peut créer un employé et activer ou désactiver les comptes. Les comptes de démonstration sont identifiables par leurs adresses et libellés ; ils doivent être remplacés avant tout usage réel.", page_before=True)
    builder.figure(SCREENSHOT_DIR / "11-admin-django.png", "Administration technique Django", "Django Admin complète l'espace de gestion métier. Il expose les modèles enregistrés aux utilisateurs autorisés, mais ne remplace pas les écrans métier où les transitions et contrôles ont été conçus pour le flux du magasin.", page_before=True)
    builder.h2("4.9 Cohérence des opérations de stock")
    builder.p("ServiceStock relit le pneu dans une transaction, calcule la nouvelle quantité, refuse un résultat négatif, met à jour le stock et crée le mouvement associé. SQLite sérialise les écritures et ne fournit pas ici de verrouillage par ligne. Cette précision est importante : l'appel select_for_update présent dans le code n'a pas l'effet qu'il aurait sur une base serveur compatible.")
    builder.p("La confirmation d'une commande appelle le service pour chaque ligne. Les sorties et le changement d'état appartiennent au même bloc transactionnel. Lors d'une annulation autorisée, les retours ne sont créés que si le stock avait été retiré et n'avait pas encore été restauré. Pour un approvisionnement, seul le statut COMMANDE permet la réception et l'opération ne peut être exécutée deux fois.")
    builder.h2("4.10 Difficultés rencontrées et corrections")
    builder.p("La première difficulté concernait la cohérence entre le panier et le stock. La solution retenue distingue la commande en attente de la confirmation par le magasin. La deuxième difficulté portait sur l'historique : les instantanés dans LigneCommande évitent qu'une modification du catalogue transforme une vente passée. La troisième concernait le parcours avant connexion, corrigé pour ne plus rediriger une requête GET vers une vue qui attend POST.")
    builder.p("Le périmètre reste centré sur le catalogue, les commandes et le suivi du stock. Les interfaces regroupent les actions utiles à chaque rôle et présentent les résultats des opérations avec des messages explicites. Les captures suivantes documentent les parcours disponibles dans la version réalisée.")

    new_body_section(doc, "Chapitre 5 - Résultats, tests et validation")
    builder.h1("Chapitre 5 - Résultats, tests et validation")
    builder.h2("5.1 Stratégie de test")
    builder.p("La suite repose sur les TestCase de Django et crée une base temporaire. Les tests de modèles vérifient les contraintes et les normalisations. Les tests de services couvrent les totaux, les transitions, les sorties, les retours et la réception fournisseur. Les tests de vues examinent les rôles, l'isolation des commandes, les pages publiques et le parcours complet du panier à la commande.")
    builder.p("Bertolino [5] rappelle que les tests répondent à des objectifs et niveaux différents. Le résultat 34 sur 34 décrit donc les scénarios exécutés. Il ne correspond ni à un pourcentage de couverture, ni à un audit de sécurité, ni à un essai de charge. L'inspection des captures complète les tests automatiques pour la mise en page visible.")
    builder.code(['python manage.py check', 'python manage.py makemigrations --check --dry-run', 'python manage.py test --verbosity 2'])
    builder.h2("5.2 Résultats observés")
    builder.table("Résultats des contrôles", ["Contrôle", "Résultat", "Portée"], [["Suite Django", "34 tests réussis, 0 échec", "Scénarios de modèles, services, vues et permissions"], ["System check", "Aucune anomalie signalée", "Configuration analysée par Django"], ["Migrations", "Aucun changement détecté", "Modèles synchronisés avec les migrations"], ["SQLite", "integrity_check = ok", "Intégrité structurelle du fichier"], ["Captures", "11 pages HTTP 200", "Visiteur, client, employé, dashboard et administrateur"]], [1.35, 1.75, 3.2])
    builder.table("Jeu de données présent dans SQLite", ["Entité", "Nombre"], [[name, value] for name, value in stats.items()], [3.8, 2.5])
    builder.h2("5.3 Validation par rapport aux objectifs")
    builder.table("Validation des objectifs", ["Objectif", "État", "Évidence"], [["Catalogue structuré et filtrable", "Atteint", "Captures publiques et tests de rendu"], ["Panier et commandes isolés par client", "Atteint", "Tests orders et captures client"], ["Stock traçable", "Atteint", "Mouvements avant et après chaque opération"], ["Approvisionnements contrôlés", "Atteint", "Réception unique au statut COMMANDE"], ["Administration par rôle", "Atteint", "Décorateurs et écrans séparés"], ["Prévision ou EOQ", "Hors périmètre", "Perspective documentée"]], [2.35, 1.1, 2.95])
    builder.h2("5.4 Limites actuelles")
    builder.p("Le serveur runserver, la clé secrète par défaut et le mode DEBUG conviennent à la démonstration locale. Ils ne forment pas une configuration de production. Les courriels sont écrits dans la console. Aucun paiement, transporteur, API WhatsApp ou service cartographique n'est intégré. Le bouton WhatsApp construit uniquement un lien wa.me.")
    builder.p("SQLite simplifie le transport du projet, mais sa concurrence d'écriture est différente de celle d'une base serveur. La suite ne contient ni test multiconnexion, ni campagne de performance, ni test de montée en charge. Une évolution vers PostgreSQL devrait être accompagnée de tests adaptés et d'une stratégie de numérotation qui ne repose pas sur count plus un.")

    new_body_section(doc, "Conclusion générale et perspectives")
    builder.h1("Conclusion générale et perspectives")
    builder.p("MANUGOM Stock répond au besoin défini par une application Django qui relie catalogue, comptes, commandes, approvisionnements et mouvements de stock. La séparation des rôles limite les accès, tandis que les services regroupent les transitions qui doivent préserver la cohérence métier. Les diagrammes et les captures montrent la même version que celle soumise aux tests.")
    builder.p("Le projet a surtout permis de traiter des problèmes concrets : conserver un prix historique, distinguer une commande en attente d'une confirmation, restaurer une quantité une seule fois et expliquer chaque variation. Les résultats confirment le fonctionnement des scénarios couverts, avec les limites propres à une démonstration locale et à une base SQLite.")
    builder.p("Les perspectives les plus utiles concernent un déploiement avec HTTPS et secrets séparés, une base serveur lorsque la concurrence l'exige, des notifications autorisées, la lecture de codes-barres, un inventaire cyclique et des tests de navigateur. Après collecte de données fiables sur la demande et les coûts, une aide au réapprovisionnement pourrait également étudier des quantités économiques, sans confondre un simple seuil avec une optimisation.")

    new_body_section(doc, "Bibliographie")
    builder.h1("Bibliographie")
    references = [
        "[1] Codd, E. F. (1970). A Relational Model of Data for Large Shared Data Banks. Communications of the ACM, 13(6), 377-387. https://doi.org/10.1145/362384.362685",
        "[2] Parnas, D. L. (1972). On the Criteria To Be Used in Decomposing Systems into Modules. Communications of the ACM, 15(12), 1053-1058. https://doi.org/10.1145/361598.361623",
        "[3] Härder, T., et Reuter, A. (1983). Principles of Transaction-Oriented Database Recovery. ACM Computing Surveys, 15(4), 287-317. https://doi.org/10.1145/289.291",
        "[4] Sandhu, R., Ferraiolo, D. F., et Kuhn, D. R. (2000). The NIST Model for Role-Based Access Control: Towards a Unified Standard. Proceedings of the Fifth ACM Workshop on Role-Based Access Control, 47-63. https://doi.org/10.1145/344287.344301",
        "[5] Bertolino, A. (2007). Software Testing Research: Achievements, Challenges, Dreams. Future of Software Engineering, 85-103. https://doi.org/10.1109/FOSE.2007.25",
        "[6] Harris, F. W. (1990). How Many Parts to Make at Once. Operations Research, 38(6), 947-950. Réimpression de l'article de 1913. https://doi.org/10.1287/opre.38.6.947",
        "[7] Django Software Foundation. Documentation Django 6.0 et 6.1. https://docs.djangoproject.com/",
        "[8] SQLite Consortium. Isolation In SQLite et documentation générale. https://sqlite.org/isolation.html",
        "[9] Bootstrap Team. Documentation Bootstrap 5.3. https://getbootstrap.com/docs/5.3/",
        "[10] Chart.js Contributors. Documentation Chart.js. https://www.chartjs.org/docs/latest/",
        "[11] StarUML. Documentation officielle. https://docs.staruml.io/",
    ]
    for ref in references:
        p = builder.p(ref)
        p.paragraph_format.left_indent = Cm(0.7); p.paragraph_format.first_line_indent = Cm(-0.7)

    new_body_section(doc, "Annexes")
    builder.h1("Annexes")
    builder.h2("Annexe A - Installation et lancement")
    builder.p("Depuis PowerShell, placer le terminal dans le dossier manugom-stock puis exécuter les commandes suivantes.")
    builder.code(['python -m venv .venv', r'.\.venv\Scripts\Activate.ps1', 'python -m pip install -r requirements.txt', 'python manage.py migrate', 'python manage.py seed_demo', 'python manage.py check', 'python manage.py test', 'python manage.py runserver'])
    builder.p("Ouvrir ensuite http://127.0.0.1:8000/. L'administration Django se trouve à l'adresse http://127.0.0.1:8000/django-admin/.")
    builder.h2("Annexe B - Comptes de démonstration")
    builder.table("Comptes fournis par seed_demo", ["Rôle", "Identifiant", "Mot de passe", "Accès"], [["Administrateur", "admin", "Admin@2026!", "Gestion et Django Admin"], ["Employé", "employe1 ou employe2", "Employe@2026!", "Espace de gestion"], ["Client", "client1 à client10", "Client@2026!", "Catalogue, panier et commandes"]], [1.1, 1.65, 1.45, 2.1])
    builder.p("Ces identifiants appartiennent au jeu de démonstration et doivent être remplacés avant un usage réel.")
    builder.h2("Annexe C - Variables de configuration")
    builder.code(['$env:DJANGO_SECRET_KEY = "votre-cle-longue-et-aleatoire"', '$env:DJANGO_DEBUG = "False"', '$env:DJANGO_ALLOWED_HOSTS = "votre-domaine.example"'])
    annex_d = builder.h2("Annexe D - Transitions principales")
    annex_d.paragraph_format.page_break_before = True
    builder.bullets(["Commande : EN_ATTENTE vers CONFIRMEE ou ANNULEE.", "Commande confirmée : EN_PREPARATION ou ANNULEE avec retour conditionnel.", "Préparation : PRETE, puis EXPEDIEE ou LIVREE selon le mode de réception.", "Approvisionnement : BROUILLON, COMMANDE, puis RECU une seule fois.", "LIVREE, ANNULEE et RECU sont des états terminaux pour les opérations correspondantes."])

    doc.save(str(OUTPUT_DOCX))
    print(f"DOCX généré : {OUTPUT_DOCX}")
    print(f"Figures : {builder.figure_number} ; tableaux : {builder.table_number}")
    return OUTPUT_DOCX


if __name__ == "__main__":
    build_docx()
